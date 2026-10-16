"""«JS-Тәлімгер» — JavaScript бойынша диагностика және ЖИ кері байланысы бар зерттеу қосымшасы."""

import csv
import hashlib
import hmac
import io
import logging
import os
import re
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (
    Flask,
    Response,
    abort,
    flash,
    redirect,
    render_template,
    request,
    session,
    url_for,
)

import ai_service
from storage import create_store
from tasks_data import ALL_TASKS, TEST_ID

log = logging.getLogger("js_talimger")

KZ_TZ = timezone(timedelta(hours=5))
CODE_RE = re.compile(r"^[A-Z0-9_-]{3,20}$")
MIN_ANSWER_LEN = 15
MAX_ANSWER_LEN = 5000
MAX_ATTEMPTS = 3  # бастапқы жауап + 2 түзету
MAX_HINTS = 2  # бір тапсырмаға ЖИ кеңесі
MAX_QUESTION_LEN = 500

AI_STATUS_LABELS = {
    "pending": "Талдау күтілуде",
    "done": "ЖИ талдауы дайын",
    "invalid": "ЖИ жауабының құрылымы бұзылған",
    "failed": "ЖИ сервисі қолжетімсіз болды",
}
REVIEW_STATUS_LABELS = {
    "not_reviewed": "Тексерілмеген",
    "confirmed": "Тексерілді: ЖИ бағасымен келісемін",
    "corrected": "Тексерілді: түзетілді",
}


def now_iso():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def make_response_id(participant_code, task_id):
    """Бір қатысушының бір тапсырмаға жауабы үшін тұрақты бірегей код (көшірмелерді болдырмайды)."""
    digest = hashlib.sha256(f"{participant_code}|{task_id}".encode()).hexdigest()
    return "R-" + digest[:8].upper()


def seed_tasks(store):
    for task in ALL_TASKS:
        existing = store.get("tasks", task["task_id"])
        if not existing or existing.get("version") != task["version"]:
            store.set("tasks", task["task_id"], task)


def create_app(store=None):
    app = Flask(__name__)
    secret = os.environ.get("SECRET_KEY")
    if not secret:
        log.warning("SECRET_KEY бапталмаған: уақытша кілт қолданылады.")
        secret = os.urandom(32)
    app.config.update(
        SECRET_KEY=secret,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_SECURE=os.environ.get("SESSION_COOKIE_SECURE", "0") == "1",
        MAX_CONTENT_LENGTH=64 * 1024,
    )
    store = store or create_store()
    app.store = store
    seed_tasks(store)
    log.warning("Деректер қоры: %s", store.name)

    @app.template_filter("kz_time")
    def kz_time(value):
        if not value:
            return "—"
        try:
            return datetime.fromisoformat(value).astimezone(KZ_TZ).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            return value

    @app.context_processor
    def inject_globals():
        return {
            "participant": session.get("participant"),
            "is_researcher": bool(session.get("researcher")),
            "AI_STATUS_LABELS": AI_STATUS_LABELS,
            "REVIEW_STATUS_LABELS": REVIEW_STATUS_LABELS,
        }

    # ---------- көмекші функциялар ----------

    def practical_tasks():
        tasks = [t for t in store.list("tasks") if t.get("type") == "practical"]
        return sorted(tasks, key=lambda t: t["task_id"])

    def get_task_or_404(task_id):
        task = store.get("tasks", task_id)
        if not task or task.get("type") != "practical":
            abort(404)
        return task

    def participant_required(view):
        @wraps(view)
        def wrapper(*args, **kwargs):
            if not session.get("participant"):
                flash("Алдымен қатысушы кодын енгізіңіз.", "warning")
                return redirect(url_for("index"))
            return view(*args, **kwargs)

        return wrapper

    def researcher_required(view):
        @wraps(view)
        def wrapper(*args, **kwargs):
            if not session.get("researcher"):
                if request.path.endswith((".csv", ".xlsx")):
                    abort(403)
                flash("Бұл бет тек зерттеушіге қолжетімді. Құпиясөзді енгізіңіз.", "warning")
                return redirect(url_for("researcher_login", next=request.path))
            return view(*args, **kwargs)

        return wrapper

    def run_ai(response_id):
        """Жауапты ЖИ-ге жібереді. Қате болса да жауап деректер қорында сақталып қалады."""
        resp = store.get("responses", response_id)
        task = store.get("tasks", resp["task_id"])
        attempts = resp.get("ai_attempts", 0) + 1
        try:
            result = ai_service.analyze_answer(task, resp["answer_text"])
        except ai_service.AIError as exc:
            log.warning("ЖИ қатесі %s: %s", response_id, exc)
            store.set(
                "responses",
                response_id,
                {"ai_status": "failed", "ai_error": str(exc), "ai_attempts": attempts, "ai_updated_at": now_iso()},
                merge=True,
            )
            return "failed"
        except Exception:
            log.exception("ЖИ талдауындағы күтпеген қате %s", response_id)
            store.set(
                "responses",
                response_id,
                {"ai_status": "failed", "ai_error": "Күтпеген қате", "ai_attempts": attempts, "ai_updated_at": now_iso()},
                merge=True,
            )
            return "failed"

        feedback = result["feedback"] or {}
        store.set(
            "ai_feedback",
            response_id,
            {
                "response_id": response_id,
                "task_id": resp["task_id"],
                "strengths": feedback.get("strengths", []),
                "errors": feedback.get("errors", []),
                "recommendation": feedback.get("recommendation", ""),
                "criteria_scores": feedback.get("criteria_scores", []),
                "suggested_score": feedback.get("suggested_score"),
                "needs_review": feedback.get("needs_review", True),
                "valid": result["status"] == "done",
                "problems": result["problems"],
                "raw_text": result["raw_text"],
                "model": result["model"],
                "prompt_version": result["prompt_version"],
                "created_at": now_iso(),
            },
        )
        store.set(
            "responses",
            response_id,
            {
                "ai_status": result["status"],
                "ai_error": "; ".join(result["problems"]),
                "ai_attempts": attempts,
                "ai_updated_at": now_iso(),
            },
            merge=True,
        )
        return result["status"]

    def can_view_response(resp):
        return session.get("researcher") or session.get("participant") == resp["participant_code"]

    def is_owner(resp):
        return session.get("participant") == resp["participant_code"]

    def attempt_no(resp):
        return resp.get("attempt", 1)

    def root_of(resp):
        return resp.get("root_id") or resp["response_id"]

    def task_attempts(root_id):
        """Бір тапсырмаға берілген барлық талпыныстар (бастапқы жауап және түзетулер)."""
        items = [store.get("responses", root_id)]
        items += store.query("responses", "root_id", root_id)
        unique = {r["response_id"]: r for r in items if r}
        return sorted(unique.values(), key=attempt_no)

    def task_hints(root_id):
        return sorted(store.query("hints", "root_id", root_id), key=lambda h: h["n"])

    def ai_score(response_id):
        fb = store.get("ai_feedback", response_id)
        return fb.get("suggested_score") if fb else None

    def answer_error(answer):
        if not answer:
            return "Жауап өрісі бос. Шешіміңізді енгізіңіз."
        if len(answer) < MIN_ANSWER_LEN:
            return f"Жауап тым қысқа. Кодыңызды және түсіндірмеңізді толығырақ жазыңыз (кемінде {MIN_ANSWER_LEN} таңба)."
        if len(answer) > MAX_ANSWER_LEN:
            return f"Жауап тым ұзын (ең көбі {MAX_ANSWER_LEN} таңба)."
        return None

    # ---------- қатысушы беттері ----------

    @app.route("/")
    def index():
        return render_template("index.html", tasks=practical_tasks())

    @app.post("/start")
    def start():
        code = request.form.get("participant_code", "").strip().upper()
        if not request.form.get("consent"):
            flash("Қатысу үшін зерттеу шарттарымен келісуіңіз қажет.", "error")
            return redirect(url_for("index"))
        if not CODE_RE.match(code):
            flash("Қатысушы коды 3–20 таңбадан тұруы тиіс: латын әріптері, сандар, «-» немесе «_».", "error")
            return redirect(url_for("index"))
        session.clear()
        session["participant"] = code
        if store.get("test_results", code):
            return redirect(url_for("tasks"))
        return redirect(url_for("test"))

    @app.route("/exit")
    def exit_participant():
        session.pop("participant", None)
        return redirect(url_for("index"))

    @app.route("/test", methods=["GET", "POST"])
    @participant_required
    def test():
        code = session["participant"]
        test_def = store.get("tasks", TEST_ID)
        if store.get("test_results", code):
            return redirect(url_for("test_result"))

        chosen = {}
        if request.method == "POST":
            missing = []
            for q in test_def["questions"]:
                value = request.form.get(q["id"])
                if value is None or not value.isdigit() or int(value) >= len(q["options"]):
                    missing.append(q["id"])
                else:
                    chosen[q["id"]] = int(value)
            if missing:
                flash(f"Барлық сұраққа жауап беріңіз (жауапсыз сұрақтар саны: {len(missing)}).", "error")
            else:
                details = [
                    {"id": q["id"], "topic": q["topic"], "chosen": chosen[q["id"]], "correct": chosen[q["id"]] == q["correct"]}
                    for q in test_def["questions"]
                ]
                store.create(
                    "test_results",
                    code,
                    {
                        "participant_code": code,
                        "test_id": TEST_ID,
                        "test_version": test_def["version"],
                        "answers": details,
                        "score": sum(d["correct"] for d in details),
                        "max_score": len(details),
                        "submitted_at": now_iso(),
                    },
                )
                return redirect(url_for("test_result"))
        return render_template("test.html", test=test_def, chosen=chosen)

    @app.route("/test/result")
    @participant_required
    def test_result():
        result = store.get("test_results", session["participant"])
        if not result:
            return redirect(url_for("test"))
        return render_template("test_result.html", result=result)

    @app.route("/tasks")
    @participant_required
    def tasks():
        code = session["participant"]
        progress = {}
        for t in practical_tasks():
            root = store.get("responses", make_response_id(code, t["task_id"]))
            if not root:
                continue
            attempts = task_attempts(root["response_id"])
            progress[t["task_id"]] = {
                "latest": attempts[-1],
                "count": len(attempts),
                "first_score": ai_score(attempts[0]["response_id"]),
                "last_score": ai_score(attempts[-1]["response_id"]),
            }
        return render_template(
            "tasks.html",
            tasks=practical_tasks(),
            progress=progress,
            test_result=store.get("test_results", code),
        )

    @app.route("/task/<task_id>", methods=["GET", "POST"])
    @participant_required
    def task(task_id):
        code = session["participant"]
        task_def = get_task_or_404(task_id)
        response_id = make_response_id(code, task_id)
        existing = store.get("responses", response_id)
        answer = ""

        if request.method == "POST":
            answer = request.form.get("answer", "").replace("\r\n", "\n").strip()
            if existing:
                flash("Бұл тапсырмаға жауабыңыз бұрын сақталған. Сақталған нәтиже көрсетілді.", "info")
                return redirect(url_for("feedback", response_id=response_id))
            error = answer_error(answer)
            if error:
                flash(error, "error")
            else:
                created = store.create(
                    "responses",
                    response_id,
                    {
                        "response_id": response_id,
                        "participant_code": code,
                        "task_id": task_id,
                        "task_version": task_def["version"],
                        "answer_text": answer,
                        "submitted_at": now_iso(),
                        "ai_status": "pending",
                        "ai_attempts": 0,
                        "attempt": 1,
                        "root_id": response_id,
                    },
                )
                if created:
                    run_ai(response_id)
                return redirect(url_for("feedback", response_id=response_id))
        elif existing:
            return redirect(url_for("feedback", response_id=response_id))

        return render_template("task.html", task=task_def, answer=answer, max_len=MAX_ANSWER_LEN)

    @app.route("/feedback/<response_id>")
    def feedback(response_id):
        resp = store.get("responses", response_id)
        if not resp:
            abort(404)
        if not can_view_response(resp):
            abort(403)
        root_id = root_of(resp)
        attempts = task_attempts(root_id)
        hints = task_hints(root_id)
        latest = attempts[-1]["response_id"] == response_id
        ready = resp.get("ai_status") in ("done", "invalid")
        return render_template(
            "feedback.html",
            resp=resp,
            task=store.get("tasks", resp["task_id"]),
            ai=store.get("ai_feedback", response_id),
            review=store.get("teacher_review", response_id),
            attempts=[dict(a, ai_score=ai_score(a["response_id"])) for a in attempts],
            hints=hints,
            owner=is_owner(resp),
            can_revise=is_owner(resp) and latest and ready and attempt_no(resp) < MAX_ATTEMPTS,
            revisions_left=MAX_ATTEMPTS - len(attempts),
            hints_left=MAX_HINTS - len(hints),
            can_hint=is_owner(resp) and latest and ready and len(hints) < MAX_HINTS,
            max_len=MAX_ANSWER_LEN,
            max_question=MAX_QUESTION_LEN,
        )

    @app.post("/feedback/<response_id>/revise")
    def revise(response_id):
        """Қатемен жұмыс: түзетілген жауапты жаңа талпыныс ретінде сақтау (бастапқы жауап өзгермейді)."""
        resp = store.get("responses", response_id)
        if not resp:
            abort(404)
        if not is_owner(resp):
            abort(403)
        root_id = root_of(resp)
        attempts = task_attempts(root_id)
        if attempts[-1]["response_id"] != response_id:
            return redirect(url_for("feedback", response_id=attempts[-1]["response_id"]))
        n = attempt_no(resp) + 1
        if n > MAX_ATTEMPTS:
            flash(f"Түзету мүмкіндігі таусылды (ең көбі {MAX_ATTEMPTS - 1} рет).", "warning")
            return redirect(url_for("feedback", response_id=response_id))
        if resp.get("ai_status") not in ("done", "invalid"):
            flash("Алдымен осы жауапқа ЖИ кері байланысын алыңыз.", "warning")
            return redirect(url_for("feedback", response_id=response_id))

        answer = request.form.get("answer", "").replace("\r\n", "\n").strip()
        error = answer_error(answer)
        if not error and answer == resp["answer_text"].strip():
            error = "Жауап өзгермеген. Кері байланысқа сүйеніп, кодты түзетіңіз."
        if error:
            flash(error, "error")
            return redirect(url_for("feedback", response_id=response_id) + "#revise")

        new_id = f"{root_id}-{n}"
        created = store.create(
            "responses",
            new_id,
            {
                "response_id": new_id,
                "participant_code": resp["participant_code"],
                "task_id": resp["task_id"],
                "task_version": resp.get("task_version"),
                "answer_text": answer,
                "submitted_at": now_iso(),
                "ai_status": "pending",
                "ai_attempts": 0,
                "attempt": n,
                "root_id": root_id,
                "parent_id": response_id,
                "hints_before": len(task_hints(root_id)),
            },
        )
        if created:
            run_ai(new_id)
        return redirect(url_for("feedback", response_id=new_id))

    @app.post("/feedback/<response_id>/hint")
    def hint(response_id):
        """ЖИ-ден шешімді айтпайтын бағыттаушы кеңес сұрау (бір тапсырмаға ең көбі MAX_HINTS)."""
        resp = store.get("responses", response_id)
        if not resp:
            abort(404)
        if not is_owner(resp):
            abort(403)
        root_id = root_of(resp)
        hints = task_hints(root_id)
        back = redirect(url_for("feedback", response_id=response_id) + "#hints")
        if len(hints) >= MAX_HINTS:
            flash(f"Кеңес саны таусылды (бір тапсырмаға {MAX_HINTS}).", "warning")
            return back
        question = request.form.get("question", "").strip()[:MAX_QUESTION_LEN]
        task_def = store.get("tasks", resp["task_id"])
        fb = store.get("ai_feedback", response_id) or {}
        try:
            model, text, version = ai_service.generate_hint(
                task_def, resp["answer_text"], fb.get("errors", []), [h["hint"] for h in hints], question
            )
        except ai_service.AIError as exc:
            log.warning("ЖИ кеңесі қатесі %s: %s", response_id, exc)
            flash("ЖИ кеңес бере алмады. Кеңес саны азайған жоқ — кейінірек қайталап көріңіз.", "warning")
            return back
        n = len(hints) + 1
        store.create(
            "hints",
            f"{root_id}-H{n}",
            {
                "hint_id": f"{root_id}-H{n}",
                "root_id": root_id,
                "response_id": response_id,
                "attempt": attempt_no(resp),
                "participant_code": resp["participant_code"],
                "task_id": resp["task_id"],
                "n": n,
                "question": question,
                "hint": text,
                "model": model,
                "prompt_version": version,
                "created_at": now_iso(),
            },
        )
        return back

    @app.post("/feedback/<response_id>/retry")
    def retry_feedback(response_id):
        resp = store.get("responses", response_id)
        if not resp:
            abort(404)
        if not can_view_response(resp):
            abort(403)
        if resp.get("ai_status") != "done":
            status = run_ai(response_id)
            if status == "failed":
                flash("ЖИ сервисі әлі қолжетімсіз. Жауабыңыз сақтаулы, кейінірек қайталап көріңіз.", "warning")
        if session.get("researcher") and request.form.get("back") == "review":
            return redirect(url_for("review", response_id=response_id))
        return redirect(url_for("feedback", response_id=response_id))

    # ---------- зерттеуші беттері ----------

    @app.route("/researcher/login", methods=["GET", "POST"])
    def researcher_login():
        password = os.environ.get("RESEARCHER_PASSWORD", "")
        if request.method == "POST":
            given = request.form.get("password", "")
            if password and hmac.compare_digest(given.encode(), password.encode()):
                session["researcher"] = True
                target = request.args.get("next", "")
                return redirect(target if target.startswith("/researcher") else url_for("researcher"))
            flash("Құпиясөз қате.", "error")
        return render_template("researcher_login.html", configured=bool(password))

    @app.route("/researcher/logout")
    def researcher_logout():
        session.pop("researcher", None)
        return redirect(url_for("index"))

    def collect_rows():
        tasks_by_id = {t["task_id"]: t for t in store.list("tasks")}
        ai_by_id = {f["response_id"]: f for f in store.list("ai_feedback")}
        reviews = {r["response_id"]: r for r in store.list("teacher_review")}
        hint_counts = {}
        for h in store.list("hints"):
            hint_counts[h["root_id"]] = hint_counts.get(h["root_id"], 0) + 1
        rows = []
        for resp in sorted(store.list("responses"), key=lambda r: r["submitted_at"]):
            rid = resp["response_id"]
            ai = ai_by_id.get(rid)
            review = reviews.get(rid)
            ai_score = ai.get("suggested_score") if ai else None
            teacher_score = review.get("teacher_score") if review else None
            diff = teacher_score - ai_score if ai_score is not None and teacher_score is not None else None
            rows.append(
                {
                    "resp": resp,
                    "task": tasks_by_id.get(resp["task_id"], {}),
                    "ai": ai,
                    "review": review,
                    "ai_score": ai_score,
                    "teacher_score": teacher_score,
                    "diff": diff,
                    "attempt": attempt_no(resp),
                    "root_id": root_of(resp),
                    "hints": hint_counts.get(root_of(resp), 0),
                }
            )
        return rows

    def compute_stats(rows, test_results):
        total = len(rows)
        ai_ok = sum(1 for r in rows if r["resp"].get("ai_status") == "done")
        reviewed = [r for r in rows if r["review"] and r["review"].get("status") in ("confirmed", "corrected")]
        compared = [r for r in reviewed if r["diff"] is not None]
        participants = {r["resp"]["participant_code"] for r in rows} | {t["participant_code"] for t in test_results}
        # Қатемен жұмыс: бастапқы және соңғы талпыныстағы ЖИ бағасының өсімі
        by_root = {}
        for r in rows:
            by_root.setdefault(r["root_id"], []).append(r)
        gains = []
        for items in by_root.values():
            items.sort(key=lambda r: r["attempt"])
            if len(items) > 1 and items[0]["ai_score"] is not None and items[-1]["ai_score"] is not None:
                gains.append(items[-1]["ai_score"] - items[0]["ai_score"])
        return {
            "revised_tasks": sum(1 for items in by_root.values() if len(items) > 1),
            "mean_gain": round(sum(gains) / len(gains), 2) if gains else None,
            "hints": len(store.list("hints")),
            "participants": len(participants),
            "responses": total,
            "tests": len(test_results),
            "ai_ok": ai_ok,
            "ai_ok_pct": round(100 * ai_ok / total) if total else 0,
            "reviewed": len(reviewed),
            "corrected": sum(1 for r in reviewed if r["review"].get("status") == "corrected"),
            "mean_abs_diff": round(sum(abs(r["diff"]) for r in compared) / len(compared), 2) if compared else None,
            "exact_match": sum(1 for r in compared if r["diff"] == 0),
            "compared": len(compared),
        }

    @app.route("/researcher")
    @researcher_required
    def researcher():
        rows = collect_rows()
        test_results = sorted(store.list("test_results"), key=lambda t: t["submitted_at"])
        test_def = store.get("tasks", TEST_ID)
        topic_errors = {}
        for t in test_results:
            for a in t["answers"]:
                if not a["correct"]:
                    topic_errors[a["topic"]] = topic_errors.get(a["topic"], 0) + 1
        return render_template(
            "researcher.html",
            rows=list(reversed(rows)),
            stats=compute_stats(rows, test_results),
            test_results=list(reversed(test_results)),
            test_def=test_def,
            topic_errors=sorted(topic_errors.items(), key=lambda kv: -kv[1]),
            store_name=store.name,
        )

    @app.route("/researcher/ai-check", methods=["GET", "POST"])
    @researcher_required
    def ai_check():
        """ЖИ тізбегіндегі әр модельдің қолжетімділігін тексеру (тек батырма басылғанда сұраныс жіберіледі)."""
        results = ai_service.probe() if request.method == "POST" else None
        chain = [("Gemini", m) for m in ai_service.gemini_models()]
        backup_on = bool(os.environ.get("BACKUP_AI_KEY") or os.environ.get("GROQ_API_KEY"))
        if backup_on:
            chain += [("Groq", m) for m in ai_service.backup_models()]
        return render_template("ai_check.html", results=results, chain=chain, backup_on=backup_on)

    @app.route("/researcher/response/<response_id>", methods=["GET", "POST"])
    @researcher_required
    def review(response_id):
        resp = store.get("responses", response_id)
        if not resp:
            abort(404)
        task_def = store.get("tasks", resp["task_id"])
        ai = store.get("ai_feedback", response_id)
        existing = store.get("teacher_review", response_id)

        if request.method == "POST":
            status = request.form.get("status", "not_reviewed")
            if status not in REVIEW_STATUS_LABELS:
                status = "not_reviewed"
            raw_score = request.form.get("teacher_score", "").strip()
            teacher_score = None
            error = None
            if raw_score:
                if raw_score.isdigit() and 0 <= int(raw_score) <= task_def["max_score"]:
                    teacher_score = int(raw_score)
                else:
                    error = f"Балл 0-ден {task_def['max_score']}-ге дейінгі бүтін сан болуы тиіс."
            if status != "not_reviewed" and teacher_score is None and not error:
                error = "Тексеру мәртебесін белгілеу үшін педагог бағасын енгізіңіз."
            if error:
                flash(error, "error")
            else:
                store.set(
                    "teacher_review",
                    response_id,
                    {
                        "response_id": response_id,
                        "teacher_score": teacher_score,
                        "teacher_comment": request.form.get("teacher_comment", "").strip(),
                        "corrected_feedback": request.form.get("corrected_feedback", "").strip(),
                        "diff_reason": request.form.get("diff_reason", "").strip(),
                        "status": status,
                        "ai_score_at_review": ai.get("suggested_score") if ai else None,
                        "reviewed_at": now_iso(),
                    },
                )
                flash("Педагог тексеруі сақталды.", "success")
                return redirect(url_for("review", response_id=response_id))

        attempts = [dict(a, ai_score=ai_score(a["response_id"])) for a in task_attempts(root_of(resp))]
        return render_template("review.html", resp=resp, task=task_def, ai=ai, review=existing,
                               attempts=attempts, hints=task_hints(root_of(resp)))

    def local_time(value):
        """ISO (UTC) уақытын Алматы уақытына (UTC+5) айналдырады: 2026-10-15 16:52:36."""
        if not value:
            return ""
        try:
            return datetime.fromisoformat(value).astimezone(KZ_TZ).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return value

    def flag(value):
        return "" if value in ("", None) else int(bool(value))

    def responses_table():
        header = [
            "response_id", "root_id", "attempt", "hints_for_task", "participant_code", "task_id", "task_version",
            "submitted_at_kz", "answer_text",
            "ai_status", "ai_model", "prompt_version", "ai_suggested_score", "max_score", "ai_strengths",
            "ai_errors", "ai_recommendation", "ai_needs_review", "teacher_status", "teacher_score",
            "score_diff_teacher_minus_ai", "teacher_comment", "corrected_feedback", "diff_reason", "reviewed_at_kz",
        ]
        out = []
        for r in collect_rows():
            resp, ai, rev = r["resp"], r["ai"] or {}, r["review"] or {}
            out.append([
                resp["response_id"], r["root_id"], r["attempt"], r["hints"], resp["participant_code"],
                resp["task_id"], resp.get("task_version"), local_time(resp["submitted_at"]), resp["answer_text"], resp.get("ai_status"), ai.get("model", ""),
                ai.get("prompt_version", ""), r["ai_score"] if r["ai_score"] is not None else "",
                r["task"].get("max_score", ""), " | ".join(ai.get("strengths", [])),
                " | ".join(ai.get("errors", [])), ai.get("recommendation", ""),
                flag(ai.get("needs_review", "")), rev.get("status", "not_reviewed"),
                r["teacher_score"] if r["teacher_score"] is not None else "",
                r["diff"] if r["diff"] is not None else "", rev.get("teacher_comment", ""),
                rev.get("corrected_feedback", ""), rev.get("diff_reason", ""), local_time(rev.get("reviewed_at", "")),
            ])
        return header, out

    def tests_table():
        test_def = store.get("tasks", TEST_ID)
        qids = [q["id"] for q in test_def["questions"]]
        header = ["participant_code", "test_version", "submitted_at_kz", "score", "max_score"]
        header += [f"{q}_correct" for q in qids]
        out = []
        for t in sorted(store.list("test_results"), key=lambda t: t["submitted_at"]):
            by_id = {a["id"]: a for a in t["answers"]}
            out.append(
                [t["participant_code"], t["test_version"], local_time(t["submitted_at"]), t["score"], t["max_score"]]
                + [int(by_id[q]["correct"]) if q in by_id else "" for q in qids]
            )
        return header, out

    def hints_table():
        header = ["hint_id", "root_id", "response_id", "attempt", "participant_code", "task_id", "n",
                  "question", "hint", "model", "prompt_version", "created_at_kz"]
        out = []
        for h in sorted(store.list("hints"), key=lambda h: h["created_at"]):
            out.append([h["hint_id"], h["root_id"], h["response_id"], h["attempt"], h["participant_code"],
                        h["task_id"], h["n"], h.get("question", ""), h["hint"], h.get("model", ""),
                        h.get("prompt_version", ""), local_time(h["created_at"])])
        return header, out

    def csv_response(filename, header, rows):
        # Орыс/қазақ тіліндегі Excel «;» бөлгішін күтеді; BOM қазақ әріптерінің дұрыс ашылуы үшін
        buf = io.StringIO()
        buf.write("﻿")
        writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
        writer.writerow(header)
        writer.writerows(rows)
        return Response(
            buf.getvalue().encode("utf-8"),
            mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    def xlsx_response(filename, sheet_title, header, rows):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title
        ws.append(header)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCE3F0")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        long_text = {"answer_text", "ai_strengths", "ai_errors", "ai_recommendation", "teacher_comment",
                     "corrected_feedback", "diff_reason", "question", "hint"}
        for idx, name in enumerate(header, start=1):
            letter = ws.cell(row=1, column=idx).column_letter
            if name in long_text:
                ws.column_dimensions[letter].width = 50
                for cell in ws[letter][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                width = max(len(str(c.value or "")) for c in ws[letter])
                ws.column_dimensions[letter].width = min(max(width + 2, 8), 28)
        buf = io.BytesIO()
        wb.save(buf)
        return Response(
            buf.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    @app.route("/researcher/export/<name>.<fmt>")
    @researcher_required
    def export(name, fmt):
        tables = {"responses": ("Жауаптар", responses_table), "tests": ("Тест нәтижелері", tests_table),
                  "hints": ("ЖИ кеңестері", hints_table)}
        if name not in tables or fmt not in ("csv", "xlsx"):
            abort(404)
        title, build = tables[name]
        header, rows = build()
        if fmt == "csv":
            return csv_response(f"{name}.csv", header, rows)
        return xlsx_response(f"{name}.xlsx", title, header, rows)

    @app.errorhandler(403)
    def forbidden(_):
        return render_template("error.html", code=403, message="Бұл бетке кіруге рұқсатыңыз жоқ."), 403

    @app.errorhandler(404)
    def not_found(_):
        return render_template("error.html", code=404, message="Бет табылмады."), 404

    return app


if __name__ == "__main__":
    from envfile import load_env_file

    load_env_file()
    logging.basicConfig(level=logging.INFO)
    create_app().run(debug=True, port=int(os.environ.get("PORT", 5000)))
