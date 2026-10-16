"""Әдістемелік нұсқаулықтағы техникалық сынақ жағдайларын тексеретін автотесттер.

Іске қосу:  python -m pytest -q
Тесттер нақты Firebase пен Gemini-ді қолданбайды: ДҚ — жадтағы LocalStore, ЖИ — алмастырылған функция.
"""

import json
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import ai_service  # noqa: E402
from app import create_app, make_response_id  # noqa: E402
from storage import LocalStore  # noqa: E402

GOOD_ANSWER = "function getEvenNumbers(numbers) {\n  return numbers.filter(n => n % 2 === 0);\n}\nfilter жұп сандарды іріктейді."

VALID_AI = {
    "strengths": ["filter әдісі дұрыс қолданылған"],
    "errors": [],
    "recommendation": "Шешімді бос массивпен тексеріп көріңіз.",
    "criteria_scores": [{"criterion": "x", "score": 2, "comment": ""}] * 4,
    "suggested_score": 8,
    "needs_review": False,
}


@pytest.fixture
def ai_calls(monkeypatch):
    """Gemini шақыруын алмастырады. calls["reply"] — ЖИ қайтаратын мәтін немесе Exception."""
    calls = {"count": 0, "reply": json.dumps(VALID_AI, ensure_ascii=False), "prompts": [],
             "hint_reply": json.dumps({"hint": "Бастапқы мәнді теріс сандармен тексеріп көріңіз."}, ensure_ascii=False),
             "hint_prompts": []}

    def fake_call(prompt, system=None, schema=None, temperature=0.2):
        if schema is ai_service.SELECT_SCHEMA:
            calls.setdefault("select_prompts", []).append(prompt)
            if isinstance(calls.get("select_reply"), Exception):
                raise calls["select_reply"]
            return "gemini-test", calls.get("select_reply") or "{}"
        if schema is ai_service.HINT_SCHEMA:
            calls["hint_prompts"].append(prompt)
            if isinstance(calls["hint_reply"], Exception):
                raise calls["hint_reply"]
            return "gemini-test", calls["hint_reply"]
        calls["count"] += 1
        calls["prompts"].append(prompt)
        if isinstance(calls["reply"], Exception):
            raise calls["reply"]
        return "gemini-test", calls["reply"]

    monkeypatch.setenv("AI_PROVIDER", "gemini")
    for name in ("GROQ_API_KEY", "BACKUP_AI_KEY", "AI_FORCE_BACKUP"):
        monkeypatch.delenv(name, raising=False)  # тесттер нақты қосалқы ЖИ-ге шықпауы тиіс
    monkeypatch.setattr(ai_service, "_call_gemini", fake_call)
    return calls


@pytest.fixture
def app(monkeypatch):
    monkeypatch.setenv("RESEARCHER_PASSWORD", "test-pass")
    monkeypatch.setenv("SECRET_KEY", "test-secret")
    application = create_app(store=LocalStore(path=None))
    application.config["TESTING"] = True
    return application


@pytest.fixture
def client(app):
    return app.test_client()


def login_participant(client, code="JS-001", tasks=("P1", "P2", "P3"), pretest=True):
    """Қатысушыны кіргізеді; pretest=True болса, алдын ала тест нәтижесі мен тапсырма тағайындауын дайындайды."""
    resp = client.post("/start", data={"participant_code": code, "consent": "1"})
    store = client.application.store
    if pretest and not store.get("test_results", code):
        answers = [{"id": f"q{i}", "topic": f"T{i}", "chosen": 0, "correct": i <= 6} for i in range(1, 11)]
        store.create("test_results", code, {"participant_code": code, "test_id": "TEST1", "test_version": 1,
                                            "kind": "pre", "answers": answers, "score": 6, "max_score": 10,
                                            "level": "mid", "submitted_at": "2026-10-16T08:00:00+00:00"})
        store.create("assignments", code, {"participant_code": code, "level": "mid", "pre_score": 6,
                                           "weak_topics": [], "task_ids": list(tasks), "reason": "тест",
                                           "chosen_by": "auto", "model": "", "prompt_version": "s1",
                                           "created_at": "2026-10-16T08:00:00+00:00"})
    return resp


def submit(client, task_id="P1", answer=GOOD_ANSWER):
    return client.post(f"/task/{task_id}", data={"answer": answer}, follow_redirects=True)


def test_correct_answer_is_saved_and_feedback_shown(client, app, ai_calls):
    login_participant(client)
    page = submit(client).get_data(as_text=True)
    rid = make_response_id("JS-001", "P1")
    assert app.store.get("responses", rid)["ai_status"] == "done"
    assert app.store.get("ai_feedback", rid)["suggested_score"] == 8
    assert "filter әдісі дұрыс қолданылған" in page
    assert "ЖИ ұсынған, педагог тексермеген кері байланыс" in page


def test_prompt_contains_task_rubric_and_delimited_answer(client, ai_calls):
    login_participant(client)
    submit(client)
    prompt = ai_calls["prompts"][0]
    assert "РУБРИКА" in prompt and "getEvenNumbers" in prompt
    assert "<<<\n" + GOOD_ANSWER + "\n>>>" in prompt


def test_empty_answer_is_rejected(client, app, ai_calls):
    login_participant(client)
    page = submit(client, answer="   ").get_data(as_text=True)
    assert "Жауап өрісі бос" in page
    assert app.store.list("responses") == []
    assert ai_calls["count"] == 0


def test_api_error_keeps_answer_and_shows_message(client, app, ai_calls):
    ai_calls["reply"] = ai_service.AIError("HTTP 503")
    login_participant(client)
    page = submit(client).get_data(as_text=True)
    rid = make_response_id("JS-001", "P1")
    resp = app.store.get("responses", rid)
    assert resp["answer_text"] == GOOD_ANSWER
    assert resp["ai_status"] == "failed"
    assert "Жауабыңыз сақталды. Кері байланысты кейін алуға болады." in page

    # Сервис қалпына келгенде қайта сұрау көшірме жасамайды.
    ai_calls["reply"] = json.dumps(VALID_AI, ensure_ascii=False)
    client.post(f"/feedback/{rid}/retry")
    assert app.store.get("responses", rid)["ai_status"] == "done"
    assert len(app.store.list("responses")) == 1
    assert len(app.store.list("ai_feedback")) == 1


def test_resubmission_does_not_create_duplicates(client, app, ai_calls):
    login_participant(client)
    submit(client)
    page = submit(client, answer=GOOD_ANSWER + " өзгертілген").get_data(as_text=True)
    assert len(app.store.list("responses")) == 1
    assert ai_calls["count"] == 1
    assert "бұрын сақталған" in page


def test_invalid_ai_structure_marked_unverified(client, app, ai_calls):
    bad = dict(VALID_AI, suggested_score=15)
    ai_calls["reply"] = json.dumps(bad)
    login_participant(client)
    submit(client)
    rid = make_response_id("JS-001", "P1")
    assert app.store.get("responses", rid)["ai_status"] == "invalid"
    fb = app.store.get("ai_feedback", rid)
    assert fb["valid"] is False and fb["needs_review"] is True and fb["suggested_score"] is None


def test_non_json_ai_reply_marked_invalid(client, app, ai_calls):
    ai_calls["reply"] = "Кешіріңіз, мен көмектесе алмаймын."
    login_participant(client)
    submit(client)
    assert app.store.get("responses", make_response_id("JS-001", "P1"))["ai_status"] == "invalid"


def test_researcher_pages_require_login(client):
    assert client.get("/researcher").status_code == 302
    assert client.get("/researcher/export/responses.csv").status_code == 403
    assert client.get("/researcher/export/responses.xlsx").status_code == 403
    client.post("/researcher/login", data={"password": "wrong"})
    assert client.get("/researcher").status_code == 302


def test_participant_cannot_view_other_feedback(client, ai_calls):
    login_participant(client, "JS-001")
    submit(client)
    rid = make_response_id("JS-001", "P1")
    login_participant(client, "JS-002")
    assert client.get(f"/feedback/{rid}").status_code == 403


def test_teacher_review_stored_separately_and_exported(client, app, ai_calls):
    login_participant(client)
    submit(client)
    rid = make_response_id("JS-001", "P1")
    client.post("/researcher/login", data={"password": "test-pass"})
    client.post(
        f"/researcher/response/{rid}",
        data={"teacher_score": "6", "status": "corrected", "diff_reason": "Түсіндірме толық емес"},
    )
    assert app.store.get("ai_feedback", rid)["suggested_score"] == 8  # ЖИ нәтижесі өзгермеді
    assert app.store.get("teacher_review", rid)["teacher_score"] == 6

    # CSV: Excel (орыс/қазақ локалі) үшін «;» бөлгіші, BOM, бір жауап — бір жол
    import csv
    import io

    csv_text = client.get("/researcher/export/responses.csv").get_data(as_text=True)
    assert csv_text.startswith("﻿response_id;")
    rows = list(csv.reader(io.StringIO(csv_text.lstrip("﻿")), delimiter=";"))
    header, row = rows[0], rows[1]
    assert len(rows) == 2 and len(row) == len(header)
    data = dict(zip(header, row))
    assert data["response_id"] == rid and data["answer_text"] == GOOD_ANSWER
    assert data["score_diff_teacher_minus_ai"] == "-2" and data["diff_reason"] == "Түсіндірме толық емес"
    assert "T" not in data["submitted_at_kz"]  # оқуға ыңғайлы уақыт

    # Excel файлы
    from openpyxl import load_workbook

    resp = client.get("/researcher/export/responses.xlsx")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.data)).active
    assert ws["A1"].value == "response_id" and ws["A2"].value == rid


def test_teacher_score_out_of_range_rejected(client, app, ai_calls):
    login_participant(client)
    submit(client)
    rid = make_response_id("JS-001", "P1")
    client.post("/researcher/login", data={"password": "test-pass"})
    client.post(f"/researcher/response/{rid}", data={"teacher_score": "12", "status": "corrected"})
    assert app.store.get("teacher_review", rid) is None


def test_test_scoring_and_export(client, app, ai_calls):
    login_participant(client, pretest=False)
    answers = {f"q{i}": "0" for i in range(1, 11)}
    client.post("/test", data=answers)
    result = app.store.get("test_results", "JS-001")
    assert result["score"] == 3  # q5, q7, q10 дұрыс жауабы — 0-нұсқа
    assert result["level"] == "low"
    client.post("/test", data={f"q{i}": "1" for i in range(1, 11)})
    assert app.store.get("test_results", "JS-001")["score"] == 3  # қайта тапсыру нәтижені өзгертпейді

    client.post("/researcher/login", data={"password": "test-pass"})
    csv_text = client.get("/researcher/export/tests.csv").get_data(as_text=True)
    assert "JS-001" in csv_text


def test_participant_code_validation(client):
    page = client.post(
        "/start", data={"participant_code": "Аслан Ахметов", "consent": "1"}, follow_redirects=True
    ).get_data(as_text=True)
    assert "Қатысушы коды" in page
    with client.session_transaction() as s:
        assert "participant" not in s


def test_ai_check_page_requires_researcher_and_runs_probe(client, monkeypatch):
    assert client.get("/researcher/ai-check").status_code == 302
    client.post("/researcher/login", data={"password": "test-pass"})
    page = client.get("/researcher/ai-check").get_data(as_text=True)
    assert "GROQ_API_KEY бапталмаған" in page
    monkeypatch.setattr(ai_service, "probe", lambda: [{"provider": "Groq", "model": "m", "ok": True, "detail": "жауап берді", "sec": 0.5}])
    page = client.post("/researcher/ai-check").get_data(as_text=True)
    assert "жауап берді" in page
